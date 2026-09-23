"""Markdown -> .docx for the VHTD manuscript (pandoc is not installed here).

    python tools/md2docx.py PAPER_v5_20260914.md          # writes PAPER_v5_20260914.docx beside it

Handles: # / ## / ### headings, paragraphs (continuation lines joined), pipe
tables, - and 1. lists, **bold** / *italic* / `code` runs, two-space-indented
display equations, and the "## Figures" list: every "- Fig. N — ... (`figNN_...png`)"
line gives the caption and the file. FIGURES ARE PLACED IN THE BODY: each one is
inserted, with its caption, right after the paragraph (or table) that first cites
"Fig. N" / "Figs. N–M"; anything never cited falls back to the Figures list at the
end (2026-09-14 evening: the author wants them inline, not collected at the end).
"""
import re, os, sys
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


def add_runs(p, text):
    pos = 0
    for m in re.finditer(r'(\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`)', text):
        if m.start() > pos:
            p.add_run(text[pos:m.start()])
        if m.group(2):
            r = p.add_run(m.group(2)); r.bold = True
        elif m.group(3):
            r = p.add_run(m.group(3)); r.italic = True
        else:
            r = p.add_run(m.group(4)); r.font.name = 'Consolas'
        pos = m.end()
    if pos < len(text):
        p.add_run(text[pos:])


def build(md_path):
    root = os.path.dirname(os.path.abspath(md_path))
    figdir = os.path.join(root, 'paper_figs')
    src = open(md_path, encoding='utf-8').read().splitlines()
    # --- pre-parse the Figures list: number -> (caption text, png or None) ---
    figs = {}
    in_list = False
    for line in src:
        if line.startswith('## '):
            in_list = (line[3:].strip() == 'Figures'); continue
        if in_list and line.startswith('- Fig. '):
            m = re.match(r'- Fig\. (\d+)( \(continued\))? — (.*)$', line)
            if not m: continue
            n = int(m.group(1)); cap = re.sub(r' \(`[^`]+`\)', '', m.group(3)).strip()
            png = re.search(r'`([^`]+\.png)`', line)
            figs.setdefault(n, []).append(('Fig. %d%s — %s' % (n, m.group(2) or '', cap),
                                           png.group(1) if png else None))
    placed = set(); seen_figs = [False]

    def insert_picture(path, cap=None):
        """Fig. 1 (fig01_block.png) is a full-page LANDSCAPE figure: it gets its own
        landscape section (11 x 8.5 in, 0.6 in margins, picture 9.6 in wide, caption
        under it) and the document returns to portrait afterwards. Every other figure
        is placed inline at 6.3 in exactly as before."""
        from docx.enum.section import WD_ORIENT, WD_SECTION
        landscape = os.path.basename(path) == 'fig01_block.png'
        if landscape:
            sec = doc.add_section(WD_SECTION.NEW_PAGE)
            sec.orientation = WD_ORIENT.LANDSCAPE
            sec.page_width, sec.page_height = Inches(11), Inches(8.5)
            for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
                setattr(sec, side, Inches(0.6))
        doc.add_picture(path, width=Inches(9.6) if landscape else Inches(6.3))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if cap:
            cp = doc.add_paragraph(); add_runs(cp, cap)
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in cp.runs: run.font.size = Pt(9)
        if landscape:
            sec = doc.add_section(WD_SECTION.NEW_PAGE)
            sec.orientation = WD_ORIENT.PORTRAIT
            sec.page_width, sec.page_height = Inches(8.5), Inches(11)
            for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
                setattr(sec, side, Inches(1.0))

    def cited(text):
        """figure numbers cited in this text, in order: Fig. 7, Figs. 7–10, 12–14, Fig. 14(a)"""
        out = []
        for m in re.finditer(r'Figs?\. ([\d–\-, ]+)', text):
            for part in re.split(r',\s*', m.group(1).strip(' ,')):
                part = part.strip()
                if not part: continue
                r = re.split(r'[–\-]', part)
                try:
                    a = int(r[0]); b = int(r[1]) if len(r) > 1 and r[1] else a
                except ValueError:
                    continue
                out.extend(range(a, b + 1))
        return out

    def place_after(text):
        if seen_figs[0]: return
        for n in cited(text):
            if n in placed or n not in figs: continue
            for cap, png in figs[n]:
                if png and os.path.exists(os.path.join(figdir, png)):
                    insert_picture(os.path.join(figdir, png), cap)
            placed.add(n)

    doc = Document()
    st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)
    i = 0; in_figs = False
    while i < len(src):
        line = src[i]
        if line.startswith('# '):
            doc.add_heading(line[2:].strip(), level=0); i += 1; continue
        if line.startswith('## '):
            h = line[3:].strip(); doc.add_heading(h, level=1); in_figs = (h == 'Figures')
            if in_figs: seen_figs[0] = True
            i += 1; continue
        if line.startswith('### '):
            doc.add_heading(line[4:].strip(), level=2); i += 1; continue
        if line.strip() == '---' or not line.strip():
            i += 1; continue
        if line.startswith('|'):
            rows = []
            while i < len(src) and src[i].startswith('|'):
                cells = [c.strip() for c in src[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'-+', c) for c in cells):
                    rows.append(cells)
                i += 1
            ncol = max(len(r) for r in rows)
            tbl = doc.add_table(rows=len(rows), cols=ncol); tbl.style = 'Table Grid'
            for ri, r in enumerate(rows):
                for ci in range(ncol):
                    cell = tbl.cell(ri, ci); cell.text = ''
                    p = cell.paragraphs[0]; add_runs(p, r[ci] if ci < len(r) else '')
                    for run in p.runs:
                        run.font.size = Pt(8)
                        if ri == 0:
                            run.bold = True
            doc.add_paragraph()
            if not in_figs:
                place_after(' '.join(' '.join(r) for r in rows))
            continue
        if line.startswith('- ') or re.match(r'^\d+\. ', line):
            style = 'List Bullet' if line.startswith('- ') else 'List Number'
            text = re.sub(r'^(- |\d+\. )', '', line)
            p = doc.add_paragraph(style=style); add_runs(p, text)
            if not in_figs:
                place_after(text)
            if in_figs:
                m = re.search(r'`([^`]+\.png)`', line)
                mn = re.match(r'- Fig\. (\d+)', line)
                if mn and int(mn.group(1)) in placed:
                    m = None
                if m and os.path.exists(os.path.join(figdir, m.group(1))):
                    insert_picture(os.path.join(figdir, m.group(1)))
                    if mn: placed.add(int(mn.group(1)))
            i += 1; continue
        if line.startswith('  ') and ('=' in line or '←' in line):
            p = doc.add_paragraph(); r = p.add_run(line.strip()); r.font.name = 'Cambria Math'
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER; i += 1; continue
        buf = [line]; i += 1
        while i < len(src) and src[i].strip() and not src[i].startswith(('#', '|', '- ', '---', '  ')) \
                and not re.match(r'^\d+\. ', src[i]):
            buf.append(src[i]); i += 1
        body = ' '.join(b.strip() for b in buf)
        p = doc.add_paragraph(); add_runs(p, body)
        if not in_figs:
            place_after(body)
    out = os.path.splitext(md_path)[0] + '.docx'
    doc.save(out)
    print('wrote', out, os.path.getsize(out), 'bytes')
    # A time-stamped READ COPY for the author, so the working file is never the one
    # open in Word (an open .docx locks the rebuild - 2026-09-14 evening).
    import shutil, datetime
    rdir = os.path.join(os.path.dirname(out), 'review'); os.makedirs(rdir, exist_ok=True)
    stem = os.path.splitext(os.path.basename(out))[0]
    rc = os.path.join(rdir, '%s_%s.docx' % (stem, datetime.datetime.now().strftime('%H%M%S')))
    shutil.copy2(out, rc); print('read copy ->', rc)
    return out


if __name__ == '__main__':
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    build(sys.argv[1])
